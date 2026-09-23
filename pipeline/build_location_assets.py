#!/usr/bin/env python3
"""Build compact Thailand location search and district-overlay assets.

Inputs are intentionally external so large source files are not committed:
- OCHA/HDX Thailand COD-AB XLSX (ADM1-ADM3 names, P-codes, centroids)
- geoBoundaries Thailand ADM2 simplified GeoJSON
- Thailand Geography JSON subdistrict file for postal-code enrichment

The official district centroid creates a one-to-one spatial join with the 928
ADM2 polygons. The builder fails closed if any district is missing, duplicated,
or mapped ambiguously.
"""
from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from shapely.geometry import Point, mapping, shape

DISTRICT_COUNT = 928
SUBDISTRICT_COUNT = 7425
BOUNDARY_TOLERANCE_DEGREES = 0.006
CENTER_TOLERANCE_DEGREES = 0.02


def rows(book: Any, sheet_name: str) -> list[dict[str, Any]]:
    sheet = book[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    header = next(iterator)
    return [dict(zip(header, row)) for row in iterator]


def numeric_code(value: Any) -> int:
    return int(str(value).replace("TH", ""))


def rounded(value: Any) -> Any:
    if isinstance(value, (list, tuple)):
        return [rounded(item) for item in value]
    if isinstance(value, float):
        return round(value, 5)
    return value


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser()
    result.add_argument("--gazetteer", required=True)
    result.add_argument("--adm2", required=True)
    result.add_argument("--postal-source", required=True)
    result.add_argument("--output-index", required=True)
    result.add_argument("--output-boundaries", required=True)
    return result


def main() -> None:
    args = parser().parse_args()
    book = load_workbook(args.gazetteer, read_only=True, data_only=True)
    provinces = rows(book, "tha_admin1")
    districts = rows(book, "tha_admin2")
    subdistricts = rows(book, "tha_admin3")
    if len(provinces) != 77 or len(districts) != DISTRICT_COUNT or len(subdistricts) != SUBDISTRICT_COUNT:
        raise RuntimeError(
            f"Unexpected COD counts: {len(provinces)} provinces, {len(districts)} districts, "
            f"{len(subdistricts)} subdistricts"
        )

    postal_rows = json.loads(Path(args.postal_source).read_text(encoding="utf-8"))
    postal_by_subdistrict = {int(row["subdistrictCode"]): int(row["postalCode"]) for row in postal_rows}
    postal_by_district: dict[int, int] = {}
    for row in postal_rows:
        postal_by_district.setdefault(int(row["districtCode"]), int(row["postalCode"]))

    province_tuples = [
        [numeric_code(row["adm1_pcode"]), str(row["adm1_name1"]), str(row["adm1_name"])]
        for row in provinces
    ]
    district_tuples = []
    district_by_code: dict[int, dict[str, Any]] = {}
    for row in districts:
        district_code = numeric_code(row["adm2_pcode"])
        province_code = numeric_code(row["adm1_pcode"])
        name_th = str(row["adm2_name1"])
        if district_code == 1008:
            name_th = "ป้อมปราบศัตรูพ่าย"
        normalized = {
            **row,
            "districtCode": district_code,
            "provinceCode": province_code,
            "districtNameTh": name_th,
            "districtNameEn": str(row["adm2_name"]),
        }
        district_by_code[district_code] = normalized
        district_tuples.append([
            province_code,
            district_code,
            name_th,
            str(row["adm2_name"]),
            postal_by_district.get(district_code),
            round(float(row["center_lat"]), 5),
            round(float(row["center_lon"]), 5),
        ])

    subdistrict_tuples = []
    for row in subdistricts:
        province_code = numeric_code(row["adm1_pcode"])
        district_code = numeric_code(row["adm2_pcode"])
        subdistrict_code = numeric_code(row["adm3_pcode"])
        subdistrict_tuples.append([
            province_code,
            district_code,
            subdistrict_code,
            str(row["adm3_name1"]),
            str(row["adm3_name"]),
            postal_by_subdistrict.get(subdistrict_code),
            round(float(row["center_lat"]), 5),
            round(float(row["center_lon"]), 5),
        ])

    index = {
        "schemaVersion": "1.0.0",
        "source": {
            "name": "Thailand COD-AB",
            "provider": "OCHA FISS / HDX",
            "sourceAgency": "Royal Thai Survey Department",
            "validOn": "2022-01-22",
            "reviewedOn": "2025-10-30",
            "license": "CC BY 3.0 IGO",
            "url": "https://data.humdata.org/dataset/cod-ab-tha",
        },
        "provinces": province_tuples,
        "districts": district_tuples,
        "subdistricts": subdistrict_tuples,
    }
    index_path = Path(args.output_index)
    index_path.parent.mkdir(parents=True, exist_ok=True)
    index_path.write_text(json.dumps(index, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")

    boundary_source = json.loads(Path(args.adm2).read_text(encoding="utf-8"))
    features = [
        (position, shape(feature["geometry"]), feature)
        for position, feature in enumerate(boundary_source["features"])
    ]
    assignments: list[tuple[int, int, dict[str, Any], Any]] = []
    for row in districts:
        district_code = numeric_code(row["adm2_pcode"])
        point = Point(float(row["center_lon"]), float(row["center_lat"]))
        containing = [(position, geometry, feature) for position, geometry, feature in features if geometry.covers(point)]
        if len(containing) == 1:
            position, geometry, feature = containing[0]
        elif not containing:
            distance, position, geometry, feature = min(
                (geometry.distance(point), position, geometry, feature)
                for position, geometry, feature in features
            )
            if distance > CENTER_TOLERANCE_DEGREES:
                raise RuntimeError(f"No district polygon covers {row['adm2_pcode']} ({row['adm2_name']})")
        else:
            raise RuntimeError(f"Ambiguous district polygon for {row['adm2_pcode']} ({row['adm2_name']})")
        assignments.append((district_code, position, feature, geometry))

    counts = Counter(position for _, position, _, _ in assignments)
    if len(assignments) != DISTRICT_COUNT or len(counts) != DISTRICT_COUNT or any(count != 1 for count in counts.values()):
        raise RuntimeError("District geometry mapping is not one-to-one")

    output_path = Path(args.output_boundaries)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    by_province: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for district_code, _, source_feature, geometry in assignments:
        district = district_by_code[district_code]
        simplified = geometry.simplify(BOUNDARY_TOLERANCE_DEGREES, preserve_topology=True)
        by_province[district["provinceCode"]].append({
            "type": "Feature",
            "properties": {
                "districtCode": district_code,
                "nameTh": district["districtNameTh"],
                "nameEn": district["districtNameEn"],
                "shapeID": source_feature["properties"]["shapeID"],
                "centerLat": round(float(district["center_lat"]), 5),
                "centerLon": round(float(district["center_lon"]), 5),
            },
            "geometry": rounded(mapping(simplified)),
        })

    written = 0
    province_payloads: dict[str, dict[str, Any]] = {}
    for province_code, province_features in sorted(by_province.items()):
        province_iso = f"TH-{province_code:02d}"
        province_payloads[province_iso] = {
            "type": "FeatureCollection",
            "provinceIso": province_iso,
            "features": sorted(province_features, key=lambda feature: feature["properties"]["districtCode"]),
        }
        written += len(province_features)
    if written != DISTRICT_COUNT or len(by_province) != 77:
        raise RuntimeError(f"Wrote {written} district polygons across {len(by_province)} provinces")
    output_path.write_text(json.dumps({
        "schemaVersion": "1.0.0",
        "source": {
            "geometry": "geoBoundaries Thailand ADM2",
            "codes": "Thailand COD-AB",
            "license": "CC BY 3.0 IGO",
        },
        "provinces": province_payloads,
    }, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")

    print(json.dumps({
        "ok": True,
        "provinces": len(province_tuples),
        "districts": len(district_tuples),
        "subdistricts": len(subdistrict_tuples),
        "boundaryFeatures": written,
        "boundaryCollections": len(by_province),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
